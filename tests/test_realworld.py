"""真实图纸语料回归测试。

这些用例全部来自桌面「图纸」目录里真实出现的内容形态（已做脱敏处理，
只保留结构），每一条都对应一个修过的真实 bug，防止回退。
"""
import pytest

from core.detect import detect
from handlers.verify import _neutralize, _placeholder_re


def hits(text, typ):
    return [text[s.start:s.end] for s in detect(text) if s.type == typ]


# ---------------------------------------------------------------- 电话不能跨行
def test_dimension_column_not_phone():
    """PDF 抽文后尺寸数字竖排，不能被 \\s 连成 400 电话。"""
    text = "地柜平面图\n400\n400\n400\n2320\n2400\n2000"
    assert hits(text, "PHONE") == []


def test_400_number_still_works():
    assert hits("热线 400-800-1234", "PHONE")


def test_australian_mobile():
    assert hits("Mobile: 0403 807 688", "PHONE")


# ---------------------------------------------------------------- NER 噪声
@pytest.mark.parametrize("text", ["格式 PDF 说明", "电话 TEL", "客户 DTC", "QTY 说明"])
def test_uppercase_abbrev_not_org(text):
    assert hits(text, "ORG") == []


def test_english_street_not_person():
    assert hits("Add: No.16, Lvjing Rd, Foshan, China", "PERSON") == []


def test_english_zone_not_org():
    assert hits("Hegui Industrial Zone, Heshun, Nanhai", "ORG") == []


def test_dimension_numbers_not_person():
    text = "抽屉(深300)\n18\n764\n18\n400\n18\n764\n18\n2000"
    assert hits(text, "PERSON") == []


def test_time_string_not_person():
    assert hits("下午2:49 回复", "PERSON") == []


def test_short_text_skips_ner():
    """短文本不跑 NER（既无收益又极慢）。"""
    assert detect("工厂") == []
    assert detect("免拉手1号") == []


# ---------------------------------------------------------------- 竖排表格
def test_vertical_workbadge_table():
    text = (
        "职位\n销售\n设计\n跟单\n订单服务助理\n"
        "姓名\n江莉晴\n张铭鉴\n潘嘉惠\n李姝\n"
        "电话\n13923273821 13924866575\n18028137590\n18028137579"
    )
    persons = hits(text, "PERSON")
    assert "江莉晴" in persons and "张铭鉴" in persons
    assert any("13923273821" in p for p in hits(text, "PHONE"))


# ---------------------------------------------------------------- xlsx
def test_header_detection_ignores_data_rows():
    from handlers.xlsx_h import _match_header

    assert _match_header("姓名") == "PERSON"
    assert _match_header("客户资料路径C:\\x\\y") is None
    assert _match_header("P:\\快盘3\\客户资料\\出口部") is None


def test_header_column_shape_guard():
    from handlers.xlsx_h import _looks_like

    assert _looks_like("Acme Trading Ltd.", "ORG")
    assert not _looks_like("Y105", "ORG")
    assert not _looks_like("1.8m", "ORG")
    assert not _looks_like("P:\\快盘3\\客户资料\\Jennie", "ORG")


def test_formula_refs_not_entities():
    from handlers.xlsx_h import _formula_skeleton

    for f in ("=G20*F20", "=SUM(F19:F21)", "=J28"):
        skel = _formula_skeleton(f)
        assert not [s for s in detect(skel) if s.score >= 0.9]


def test_sparse_cells_iteration():
    """max_column=16384 的表不能被全量遍历（实测 620 万格 -> 300 秒超时）。"""
    import tempfile
    from pathlib import Path

    from openpyxl import Workbook, load_workbook

    from handlers.xlsx_h import _iter_cells

    with tempfile.TemporaryDirectory() as td:
        p = Path(td) / "wide.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "姓名"
        ws["B1"] = "电话"
        ws["A2"] = "张三"
        ws["B2"] = "13800138000"
        ws["XFD9"] = None          # 撑大 max_column
        wb.save(str(p))
        wb2 = load_workbook(str(p))
        cells = list(_iter_cells(wb2.worksheets[0]))
        assert len(cells) < 100, f"应只遍历真实存在的单元格，实际 {len(cells)}"
        wb2.close()


# ---------------------------------------------------------------- 旧格式
def test_xls_reader_on_real_biff():
    """自写 BIFF 读取器：非 OLE2 / 损坏文件必须明确报错，不能静默返回空表。"""
    from handlers.xls_reader import CFBFile, OLEError

    with pytest.raises(OLEError):
        CFBFile.load(b"not an ole2 file at all")
    # 只有 OLE2 魔数、后续全零 -> 目录为空，读不到 Workbook 流
    blob = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600
    try:
        cfb = CFBFile.load(blob)
        assert cfb.stream(("Workbook", "Book")) is None
    except OLEError:
        pass


def test_html_disguised_xls_detected(tmp_path):
    from handlers.dispatch import sniff

    p = tmp_path / "fake.xls"
    p.write_bytes(b'<html xmlns:v="urn:schemas-microsoft-com:office:excel"><table><tr><td>x</td></tr></table></html>')
    assert sniff(p) == "html_xls"


# ---------------------------------------------------------------- 复核
def test_placeholders_kept_in_verify():
    """复核时占位符必须保留（它是「已脱敏」的正向证据），只中和字段标签。"""
    text = "姓名：[姓名]  电话：[手机号]  公司名称：某公司"
    out = _neutralize(text)
    assert "[姓名]" in out and "[手机号]" in out
    assert "公司名称" not in out


def test_placeholder_regex_matches_all():
    rx = _placeholder_re()
    for tok in ("[姓名]", "[公司]", "[地址]", "[手机号]", "[银行卡号]", "[证件号]", "[邮箱]"):
        assert rx.search(tok), tok


# ---------------------------------------------------------------- 临时目录
def test_workdir_uses_mkdir_not_mkdtemp():
    """mkdtemp 在受限 ACL 环境下会造出拒绝写入的目录（实测 PermissionError）。"""
    from core.workdir import temporary_workdir

    with temporary_workdir("test") as d:
        (d / "probe.txt").write_bytes(b"hello")
        assert (d / "probe.txt").read_bytes() == b"hello"
    assert not d.exists()
